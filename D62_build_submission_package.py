# -*- coding: utf-8 -*-
"""
D62_build_submission_package.py — assemble the final EJE submission package.

Rebuilds EJE/ from the corrected sources:
  two-step-MR-overlap_manuscript.docx   (from D61_论文终稿_20260906.md via D61_build_docx.py)
  Cover_letter_EJE.docx                 (from cover_letter_EJE_20260910.md, internal notes stripped)
  Figures/Fig1-4 + FigS1,S2,S3(regime map),S4  (vector PDF)
  Supplementary_material/*.docx         (7 supplements converted from the corrected .md)

Safety: the workspace delete-guard blocks rmtree on Chinese paths, so obsolete files are
MOVED out to a recycle folder (reversible) instead of deleted.

Run: python D62_build_submission_package.py  (after D61_build_docx.py)
"""
import io, os, re, shutil, subprocess, sys

ROOT = r"D:\WorkBuddy\两步法MR中介"
SRC = os.path.join(ROOT, "⑤ 投稿文件_20260904")
PKG = os.path.join(ROOT, "EJE")
RECYCLE = os.path.join(ROOT, "_退役_EJE旧件")
PY = sys.executable

# --- supplement mapping: source .md -> package .docx name ---
SUPP = [
    ("D61_Supplementary_S1-S6_derivation.md",    "Supplementary_S1-S6_Covariance_derivation.docx"),
    ("D61_Supplementary_S9_prisma_checklist.md", "Supplementary_Table_S2_PRISMA2020_checklist.docx"),
    ("D61_Supplementary_S3_batch.md",            "Supplementary_Table_S3_batch_reestimation.docx"),
    ("D61_Supplementary_S7_kappa.md",            "Supplementary_Methods_S7_intercoder_reliability.docx"),
    ("D61_Supplementary_S8_mvmr.md",             "Supplementary_Methods_S8_MVMR_generalization.docx"),
    ("D61_Supplementary_STROBE_MR.md",           "Supplementary_STROBE-MR_checklist.docx"),
    ("D61_Supplementary_codebook_S1.md",         "Supplementary_Codebook_S1.docx"),
]

MAIN_FIGS = ["Fig1_conceptual.pdf", "Fig2_PRISMA.pdf", "Fig3_literature.pdf", "Fig4_realdata.pdf"]
SUPP_FIGS = ["FigS1_simulation.pdf", "FigS2_pisweep.pdf", "FigS4_batch_forest.pdf"]

OBSOLETE_FIG = "FigS3_flip_probability.pdf"
NEW_FIG = "FigS3_regime_map.pdf"

# ---------------------------------------------------------------- md -> docx
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _inline(par, text):
    pat = re.compile(r'(\*\*([^*]+)\*\*)|(\*([^*]+)\*)|(`([^`]+)`)|(\$([^$]+)\$)')
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(2) is not None:
            r = par.add_run(m.group(2)); r.bold = True
        elif m.group(4) is not None:
            r = par.add_run(m.group(4)); r.italic = True
        elif m.group(6) is not None:
            r = par.add_run(m.group(6)); r.font.name = "Consolas"
        elif m.group(8) is not None:
            r = par.add_run(m.group(8)); r.italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def md_to_docx(src, dst):
    lines = io.open(src, encoding="utf-8").read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        if not s.strip():
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            doc.add_heading(m.group(2), level=min(lvl, 4)); i += 1; continue
        if s.strip().startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            header = [c.strip() for c in s.strip().strip("|").split("|")]
            i += 2; rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
            for j, h in enumerate(header):
                _inline(t.rows[0].cells[j].paragraphs[0], h)
                for rn in t.rows[0].cells[j].paragraphs[0].runs:
                    rn.bold = True
            for row in rows:
                cells = t.add_row().cells
                for j in range(min(len(row), len(header))):
                    _inline(cells[j].paragraphs[0], row[j])
            doc.add_paragraph(); continue
        mb = re.match(r'^[-*]\s+(.*)$', s)
        if mb:
            _inline(doc.add_paragraph(style="List Bullet"), mb.group(1)); i += 1; continue
        mn = re.match(r'^\d+[.)]\s+(.*)$', s)
        if mn:
            _inline(doc.add_paragraph(style="List Number"), mn.group(1)); i += 1; continue
        _inline(doc.add_paragraph(), s); i += 1
    doc.save(dst)


def build_cover(md_path, out_path):
    raw = io.open(md_path, encoding="utf-8", newline="").read()
    start = raw.find("Dear Editors")
    if start < 0:
        start = raw.find("Dear Professor")
    body = raw[start:]
    end = body.find("*Internal notes")
    if end > 0:
        body = body[:end]
    body = body.replace("---", "")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    for para in body.split("\n"):
        t = para.strip()
        if not t:
            continue
        p = doc.add_paragraph()
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t.strip("*")); r.bold = True
        else:
            _inline(p, t)
        for r in p.runs:
            r.font.name = "Times New Roman"; r.font.size = Pt(11)
    doc.save(out_path)


def move_out(path):
    """Move an obsolete file into the (reversible) recycle folder."""
    if not os.path.exists(path):
        return
    os.makedirs(RECYCLE, exist_ok=True)
    dst = os.path.join(RECYCLE, os.path.basename(path))
    k = 1
    while os.path.exists(dst):
        dst = os.path.join(RECYCLE, f"{k}_" + os.path.basename(path)); k += 1
    shutil.move(path, dst)
    print(f"   moved obsolete -> {dst}")


def main():
    fig_dir = os.path.join(PKG, "Figures")
    sup_dir = os.path.join(PKG, "Supplementary_material")
    os.makedirs(fig_dir, exist_ok=True)
    os.makedirs(sup_dir, exist_ok=True)

    # 1) manuscript docx (rebuild first)
    print("[1] rebuilding manuscript docx ...")
    subprocess.run([PY, os.path.join(SRC, "D61_build_docx.py")], check=True)
    ms_src = os.path.join(SRC, "D61_论文终稿_20260906.docx")
    shutil.copy2(ms_src, os.path.join(PKG, "two-step-MR-overlap_manuscript.docx"))

    # 2) cover letter
    print("[2] building cover letter ...")
    build_cover(os.path.join(SRC, "cover_letter_EJE_20260910.md"),
                os.path.join(PKG, "Cover_letter_EJE.docx"))

    # 3) figures
    print("[3] figures ...")
    for f in MAIN_FIGS + SUPP_FIGS:
        src = os.path.join(SRC, "figure", f)
        if not os.path.exists(src):
            print("   [WARN] missing figure:", f); continue
        shutil.copy2(src, os.path.join(fig_dir, f))
    move_out(os.path.join(fig_dir, OBSOLETE_FIG))          # retire the stale flip figure
    shutil.copy2(os.path.join(SRC, "figure", NEW_FIG), os.path.join(fig_dir, NEW_FIG))

    # 4) supplements: convert corrected .md -> .docx
    print("[4] supplements ...")
    for md_name, docx_name in SUPP:
        src = os.path.join(SRC, md_name)
        if not os.path.exists(src):
            print("   [WARN] missing supplement:", md_name); continue
        md_to_docx(src, os.path.join(sup_dir, docx_name))
        print("   ", docx_name)

    # 5) report
    print("\n=== FINAL PACKAGE:", PKG, "===")
    total = 0
    for r, d, files in os.walk(PKG):
        rel = os.path.relpath(r, PKG)
        for f in sorted(files):
            total += 1
            print("  ", (os.path.join(rel, f) if rel != "." else f))
    print(f"  ({total} files)")


if __name__ == "__main__":
    main()
